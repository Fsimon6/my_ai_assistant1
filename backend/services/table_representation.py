# -*- coding: utf-8 -*-
"""
通用二维表格 Representation Layer（第一阶段）

职责：
  1. 解析 .xlsx / .xls / .csv / .tsv  -> 统一中间网格（Unified Table Representation）
  2. 构建 Table-aware Chunks（Workbook Summary / Sheet Schema / Row Group）
  3. 解析后校验（§21：防止异常 Excel 静默丢列）
  4. 持久化原始文件与 Representation JSON（§18/§19）

解析核心使用 excel-parser（benchmark 选定），CSV/TSV 用内置 csv 读取器，
两者都归一化为相同的“网格”结构，后续构建逻辑与格式无关。
"""
import os
import csv
import json
import shutil
from typing import List, Dict, Any, Optional
from pathlib import Path
from datetime import datetime

# 表格类扩展名（与 document_service.ALLOWED_EXT / api/v1/rag.py 保持一致）
TABLE_EXTENSIONS = {'.xlsx', '.xls', '.csv', '.tsv'}


def col_letter(n: int) -> str:
    """1->A, 26->Z, 27->AA ..."""
    s = ''
    while n > 0:
        n, r = divmod(n - 1, 26)
        s = chr(65 + r) + s
    return s


def _is_numeric(val: Any) -> bool:
    if val is None:
        return False
    s = str(val).strip().replace(',', '').replace('%', '')
    if s == '':
        return False
    try:
        float(s)
        return True
    except ValueError:
        return False


def _is_date(val: Any) -> bool:
    if val is None:
        return False
    s = str(val)
    # 简单启发：含日期/时间分隔符的文本
    return ('/' in s or '-' in s) and any(ch.isdigit() for ch in s) and (':' in s or s.count('/') >= 2 or s.count('-') >= 2)


# ----------------------------- 解析为统一网格 -----------------------------
def _parse_with_excel_parser(file_path: str) -> List[Dict[str, Any]]:
    """用 excel-parser 读取单元格图（逐单元格扫描，对 <dimension> 损坏文件健壮）。"""
    from excel_parser import parse_workbook
    res = parse_workbook(path=file_path)
    out: List[Dict[str, Any]] = []
    for sh in res.workbook.sheets:
        cells = sh.cells or {}
        if not cells:
            continue
        max_r = max(int(k.split(',')[0]) for k in cells)
        max_c = max(int(k.split(',')[1]) for k in cells)
        rows = []
        for r in range(1, max_r + 1):
            row = []
            for c in range(1, max_c + 1):
                cell = cells.get(f"{r},{c}")
                if cell is None:
                    row.append(None)
                else:
                    val = getattr(cell, 'display_value', None)
                    if val is None:
                        val = getattr(cell, 'raw_value', None)
                    row.append({'value': val, 'data_type': getattr(cell, 'data_type', 's')})
            rows.append(row)
        out.append({
            'sheet_name': sh.sheet_name,
            'sheet_index': getattr(sh, 'sheet_index', len(out)),
            'sheet_id': getattr(sh, 'sheet_id', getattr(sh, 'sheet_index', len(out))),
            'rows': rows,
        })
    return out


def _parse_csv_tsv(file_path: str, delimiter: str) -> List[Dict[str, Any]]:
    with open(file_path, 'r', encoding='utf-8-sig', newline='') as f:
        reader = list(csv.reader(f, delimiter=delimiter))
    if not reader:
        return []
    max_c = max((len(r) for r in reader), default=0)
    rows = []
    for r in reader:
        row = [{'value': (v if v != '' else None), 'data_type': 's'} for v in r]
        while len(row) < max_c:
            row.append(None)
        rows.append(row)
    return [{'sheet_name': 'Sheet1', 'sheet_index': 0, 'sheet_id': 0, 'rows': rows}]


def parse_to_grids(file_path: str) -> List[Dict[str, Any]]:
    ext = Path(file_path).suffix.lower()
    if ext == '.csv':
        return _parse_csv_tsv(file_path, ',')
    if ext == '.tsv':
        return _parse_csv_tsv(file_path, '\t')
    return _parse_with_excel_parser(file_path)


# ----------------------------- 表头检测 -----------------------------
def _row_has_value_cells(row: List[Optional[Dict]]) -> bool:
    """判断一行是否像“真实数据行”而非表头/描述行。

    真实数据行通常含多个数值单元格（数量/金额/ID 等），或含长数字 ID
    （如 18 位 Order ID）。表头行与描述行主要是字段名/说明文本，不含这类值。

    不能用“行内是否含数字”判定数据行：标题行可能含数字（如某列名为 '50'），
    描述行也可能含数字（说明文本里出现编号），二者都不应被误判为数据。
    """
    numeric = 0
    long_id = 0
    for cell in row:
        if not cell:
            continue
        v = cell.get('value')
        if v is None or str(v).strip() == '':
            continue
        if _is_numeric(v):
            numeric += 1
        s = str(v).strip()
        if s.isdigit() and len(s) >= 10:
            long_id += 1
    return numeric >= 2 or long_id >= 1


def _detect_header_rows(rows: List[List[Optional[Dict]]]) -> int:
    """返回前导表头行数（至少 1 行）。

    第一行恒为字段名表头；其后连续的“非数据行”也视为表头，
    因此支持“标题行 + 描述行”两行表头（甚至更多层）。

    判定数据行开始的信号：该行含 >=2 个数值单元格，或含长数字 ID。
    描述行（纯说明文本）与含数字的标题行都不会触发，从而正确区分：
        标题行 / 字段名行 / 字段描述行 / 真正数据行。
    """
    hdr = 0
    for r in rows:
        if all(cell is None or (cell.get('value') in (None, '')) for cell in r):
            break  # 空行 => 数据开始
        # 第一行恒为表头；其余行若为真实数据行则表头结束
        if hdr >= 1 and _row_has_value_cells(r):
            break
        hdr += 1
        if hdr >= 4:  # 保守上限，避免把整张表误判为表头
            break
    return max(hdr, 1)


def _col_datatype(data_rows: List[Dict], c: int) -> str:
    vals = []
    for row in data_rows:
        if c < len(row) and row[c] and row[c].get('value') is not None:
            vals.append(row[c]['value'])
    if not vals:
        return 'text'
    if all(_is_numeric(v) for v in vals):
        return 'numeric'
    if all(_is_date(v) for v in vals):
        return 'date'
    return 'text'


# ----------------------------- 构建 Unified Representation -----------------------------
def build_representation(
    file_path: str,
    document_id: str,
    user_id: Any,
    filename: str,
    file_type: str,
    original_path: str,
    embedding_model: Optional[str] = None,
) -> Dict[str, Any]:
    grids = parse_to_grids(file_path)
    sheets_rep = []
    for g in grids:
        rows = g['rows']
        n_hdr = _detect_header_rows(rows)
        header_rows = rows[:n_hdr]
        data_rows = rows[n_hdr:]
        n_cols = max((len(r) for r in rows), default=0)

        columns = []
        for c in range(1, n_cols + 1):
            tech = None
            if len(header_rows[0]) > c - 1 and header_rows[0][c - 1]:
                tech = header_rows[0][c - 1].get('value')
            tech = str(tech) if tech is not None else f"Column{c}"
            desc = None
            if n_hdr >= 2 and len(header_rows[1]) > c - 1 and header_rows[1][c - 1]:
                desc = str(header_rows[1][c - 1].get('value'))
            display = tech
            if desc and len(desc) <= 40 and not desc.rstrip().endswith(('.', '。', ';', '；', ':')):
                display = desc
            columns.append({
                'col_index': c,
                'col_letter': col_letter(c),
                'technical_name': tech,
                'display_name': display,
                'description': desc,
                'data_type': _col_datatype(data_rows, c - 1),
            })

        row_rep = []
        for i, dr in enumerate(data_rows):
            abs_row = n_hdr + 1 + i
            cells = []
            for c in range(1, n_cols + 1):
                cell = dr[c - 1] if c - 1 < len(dr) else None
                cells.append({
                    'col_index': c,
                    'col_letter': col_letter(c),
                    'value': (cell.get('value') if cell else None),
                    'type': (cell.get('data_type') if cell else 's'),
                })
            rng = f"{col_letter(1)}{abs_row}:{col_letter(n_cols)}{abs_row}"
            row_rep.append({'row_index': abs_row, 'range': rng, 'cells': cells})

        table_range = f"A1:{col_letter(n_cols)}{len(rows)}"
        sheets_rep.append({
            'sheet_id': g['sheet_id'],
            'sheet_name': g['sheet_name'],
            'tables': [{
                'table_id': f"{g['sheet_name']}#0",
                'range': table_range,
                'header_rows': list(range(1, n_hdr + 1)),
                'n_header_rows': n_hdr,
                'columns': columns,
                'data_row_start': n_hdr + 1,
                'data_row_end': len(rows),
                'row_count': len(data_rows),
                'col_count': n_cols,
                'rows': row_rep,
            }]
        })

    return {
        'document_id': document_id,
        'filename': filename,
        'file_type': file_type,
        'user_id': user_id,
        'original_path': original_path,
        'embedding_model': embedding_model,
        'generated_at': datetime.now().isoformat(),
        'workbook': {'sheets': sheets_rep},
        'source_uri_template': f"{filename}#{{sheet}}!{{range}}",
    }


# ----------------------------- 解析后校验（§21） -----------------------------
class TableParseValidationError(Exception):
    pass


def validate_representation(file_path: str, rep: Dict[str, Any]) -> None:
    """独立扫描原始文件，对比 Representation 的列数，防止静默丢列。"""
    ext = Path(file_path).suffix.lower()
    raw_cols = None
    try:
        if ext in ('.xlsx', '.xls'):
            import openpyxl
            wb = openpyxl.load_workbook(file_path, read_only=True, data_only=True)
            ws = wb[wb.sheetnames[0]]
            raw_cols = ws.max_column
            wb.close()
        elif ext in ('.csv', '.tsv'):
            delim = ',' if ext == '.csv' else '\t'
            with open(file_path, encoding='utf-8-sig', newline='') as f:
                raw_cols = max((len(r) for r in csv.reader(f, delimiter=delim)), default=0)
    except Exception as e:  # 校验失败不应阻断主流程，仅记录
        import logging
        logging.getLogger(__name__).warning(f'表格校验扫描失败（跳过）：{e}')
        return

    rep_cols = max(
        (t['col_count'] for s in rep['workbook']['sheets'] for t in s['tables']),
        default=0
    )
    if raw_cols and rep_cols and rep_cols < raw_cols / 2:
        raise TableParseValidationError(
            f'解析严重丢列：原始文件 {raw_cols} 列，Representation 仅 {rep_cols} 列'
        )


# ----------------------------- 构建 Table-aware Chunks -----------------------------
def _mk_chunk(idx, chunk_type, sheet_name, table, rep, source_uri, rstart, rend, content, processed_at):
    rng = source_uri.split('!')[-1] if '!' in source_uri else table['range']
    meta = {
        'document_id': rep['document_id'],
        'user_id': rep['user_id'],
        'filename': rep['filename'],
        'file_type': rep['file_type'],
        'embedding_model': rep.get('embedding_model'),
        'chunk_type': chunk_type,
        'sheet_name': sheet_name,
        'table_id': table['table_id'],
        'row_start': rstart,
        'row_end': rend,
        'column_start': 1,
        'column_end': table['col_count'],
        'range': rng,
        'source_uri': source_uri,
        'source': source_uri,
        'processed_at': processed_at,
        'chunk_index': idx,
    }
    return {'id': f"{rep['document_id']}_{idx}", 'content': content, 'metadata': meta}


def build_table_chunks(rep: Dict[str, Any], processed_at: str) -> List[Dict[str, Any]]:
    chunks = []
    idx = 0
    filename = rep['filename']
    for s in rep['workbook']['sheets']:
        sheet_name = s['sheet_name']
        for t in s['tables']:
            col_names = [c['technical_name'] for c in t['columns']]
            # 1) Workbook / Sheet Summary
            desc_examples = "; ".join(
                f"{c['technical_name']}（{c['description']}）"
                for c in t['columns'][:12] if c['description']
            )
            summary = (
                f"文件《{filename}》（类型 {rep['file_type']}）是一个二维表格文档，"
                f"包含 {len(rep['workbook']['sheets'])} 个工作表。\n"
                f"工作表「{sheet_name}」范围 {t['range']}，共 {t['col_count']} 列、"
                f"{t['row_count']} 条数据行（含 {t['n_header_rows']} 层表头）。\n"
                f"该表全部字段（列字母=字段名）：\n"
                + ", ".join(f"{c['col_letter']}={c['technical_name']}" for c in t['columns'])
                + (f"\n部分字段描述：{desc_examples}" if desc_examples else "")
            )
            chunks.append(_mk_chunk(idx, 'workbook_summary', sheet_name, t, rep,
                                    f"{filename}#{sheet_name}", 0, 0, summary, processed_at))
            idx += 1

            # 2) Sheet Schema（字段结构，带 table_id 便于与 Row 关联）
            schema_lines = [f"工作表「{sheet_name}」字段结构（表 {t['table_id']}，共 {t['col_count']} 列）："]
            for c in t['columns']:
                line = f"{c['col_letter']} {c['technical_name']}"
                if c['description']:
                    line += f" — {c['description']}"
                line += f" （类型：{c['data_type']}）"
                schema_lines.append(line)
            hdr_range = f"A1:{col_letter(t['col_count'])}{t['n_header_rows']}"
            chunks.append(_mk_chunk(idx, 'sheet_schema', sheet_name, t, rep,
                                    f"{filename}#{sheet_name}!{hdr_range}", 1, t['n_header_rows'],
                                    "\n".join(schema_lines), processed_at))
            idx += 1

            # 3) Row Group（每 GROUP 行一组，字段映射仅写一次）
            GROUP = 8
            data_rows = t['rows']
            mapping = "字段映射（列字母=字段名）：" + ", ".join(
                f"{c['col_letter']}={c['technical_name']}" for c in t['columns']
            )
            for gi in range(0, len(data_rows), GROUP):
                group = data_rows[gi:gi + GROUP]
                rstart = group[0]['row_index']
                rend = group[-1]['row_index']
                rng = f"A{rstart}:{col_letter(t['col_count'])}{rend}"
                header_line = "行号 | " + " | ".join(c['col_letter'] for c in t['columns'])
                line_rows = []
                for row in group:
                    vals = [("" if cell['value'] is None else str(cell['value'])) for cell in row['cells']]
                    line_rows.append(f"{row['row_index']} | " + " | ".join(vals))
                text = (
                    f"工作表「{sheet_name}」| 数据范围：{rng}（第 {rstart}-{rend} 行，共 {len(group)} 行）\n"
                    f"{mapping}\n"
                    f"{header_line}\n"
                    + "\n".join(line_rows)
                )
                chunks.append(_mk_chunk(idx, 'row_group', sheet_name, t, rep,
                                        f"{filename}#{sheet_name}!{rng}", rstart, rend, text, processed_at))
                idx += 1
    return chunks


# ----------------------------- 原始文件与 Representation 持久化（§18/§19） -----------------------------
def table_store_dir() -> Path:
    # 与 Chroma（settings.VECTOR_DB_PATH 的同级 data 目录）保持一致，便于管理
    from backend.config.settings import settings
    return Path(settings.VECTOR_DB_PATH).resolve().parent / 'table_originals'


def store_original(file_path: str, document_id: str) -> str:
    d = table_store_dir()
    d.mkdir(parents=True, exist_ok=True)
    ext = Path(file_path).suffix.lower()
    dest = d / f"{document_id}{ext}"
    try:
        shutil.move(file_path, dest)
    except Exception:
        shutil.copy2(file_path, dest)
    return str(dest)


def save_representation_json(rep: Dict[str, Any]) -> str:
    d = table_store_dir()
    d.mkdir(parents=True, exist_ok=True)
    p = d / f"{rep['document_id']}.json"
    p.write_text(json.dumps(rep, ensure_ascii=False, indent=2), encoding='utf-8')
    return str(p)


def load_representation(document_id: str) -> Optional[Dict[str, Any]]:
    p = table_store_dir() / f"{document_id}.json"
    if not p.exists():
        return None
    return json.loads(p.read_text(encoding='utf-8'))
