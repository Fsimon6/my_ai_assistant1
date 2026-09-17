import openpyxl
P = r'C:\Users\Administrator\Desktop\my_ai_assistant\data\table_originals\1f1e9dc28e3f4d4faf9e140d55277b48.xlsx'
wb = openpyxl.load_workbook(P, read_only=True, data_only=True)
ws = wb[wb.sheetnames[0]]
print('sheetnames', wb.sheetnames)
print('max_column', ws.max_column, 'max_row', ws.max_row)
print('dimensions', ws.dimensions)
non_empty = set()
for row in ws.iter_rows():
    for c in row:
        if c.value is not None and str(c.value).strip() != '':
            non_empty.add(c.column)
print('non_empty_cols', len(non_empty))
# 打印前 21 行、前 15 列的表头与样本
print('--- header+sample (first 21 rows, first 15 cols) ---')
for ri, row in enumerate(ws.iter_rows(min_row=1, max_row=21, max_col=15), 1):
    vals = [c.value for c in row]
    print(ri, vals)
wb.close()
