import os, shutil, csv

ROOT = r'C:\Users\Administrator\Desktop\my_ai_assistant'
REAL_XLSX = os.path.join(ROOT, 'data', 'table_originals', '1f1e9dc28e3f4d4faf9e140d55277b48.xlsx')

# 1) 真实 XLSX（63 列订单表）复制为测试上传文件，避免改动原文件
shutil.copyfile(REAL_XLSX, os.path.join(ROOT, '_test_order.xlsx'))
print('XLSX_COPIED')

# 2) 真实 CSV
with open(os.path.join(ROOT, '_test.csv'), 'w', encoding='utf-8', newline='') as f:
    w = csv.writer(f)
    w.writerow(['product', 'price', 'quantity', 'city'])
    w.writerow(['Apple', '5.0', '10', 'Beijing'])
    w.writerow(['Banana', '3.0', '20', 'Shanghai'])
print('CSV_WRITTEN')

# 3) 真实 TSV
with open(os.path.join(ROOT, '_test.tsv'), 'w', encoding='utf-8', newline='') as f:
    f.write('col1\tcol2\tcol3\n')
    f.write('a\t1\tx\n')
    f.write('b\t2\ty\n')
print('TSV_WRITTEN')

# 4) 文本 / markdown
with open(os.path.join(ROOT, '_test.txt'), 'w', encoding='utf-8') as f:
    f.write('This is a plain text upload regression test file.')
with open(os.path.join(ROOT, '_test.md'), 'w', encoding='utf-8') as f:
    f.write('# Markdown Test\n\nSome **markdown** content for upload regression.')
print('TXT_MD_WRITTEN')

# 5) DOCX（若 python-docx 可用则生成真实 docx）
try:
    from docx import Document
    d = Document()
    d.add_heading('Docx Upload Test', 0)
    d.add_paragraph('This is a real docx test file for upload regression verification.')
    d.save(os.path.join(ROOT, '_test.docx'))
    print('DOCX_REAL_WRITTEN')
except Exception as e:
    print('DOCX_MISSING', repr(e))

# 6) XLS：用于校验白名单放行（真实解析需 xlrd，仅生成占位文件以验证“不被扩展名拒绝”）
with open(os.path.join(ROOT, '_test.xls'), 'wb') as f:
    f.write(b'DUMMY_XLS_FOR_GATE_TEST')
print('XLS_PLACEHOLDER_WRITTEN')
